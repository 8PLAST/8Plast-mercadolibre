import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from bulk_import import BAG_COLUMNS, ROLL_COLUMNS, create_template, import_preview, preview_file
from db import Database


class BulkImportTests(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory(); self.root=Path(self.folder.name)
        self.db=Database(self.root/"stock.db")

    def tearDown(self):self.folder.cleanup()

    def workbook(self,bag_rows=None,roll_rows=None):
        path=self.root/"products.xlsx"; book=Workbook(); book.remove(book.active)
        for name,headers,rows in (("Bolsas",BAG_COLUMNS,bag_rows or []),("Rollos",ROLL_COLUMNS,roll_rows or [])):
            sheet=book.create_sheet(name); sheet.append(headers)
            for row in rows:sheet.append(row)
        book.save(path); return path

    def valid_bag(self,sku="B-100",initial=1000):
        return [sku,"Bolsa prueba",50,70,30,"Transparente","Cortada",50,100,2000,initial,"SI"]

    def valid_roll(self,sku="R-100",initial=20):
        return [sku,"Rollo prueba",30,50,100,"Transparente","Tubo",5,20,initial,"SI"]

    def test_valid_import_and_initial_stock_movement(self):
        rows=preview_file(self.workbook([self.valid_bag()],[self.valid_roll()]),self.db)
        self.assertEqual([row.status for row in rows],["VÁLIDO","VÁLIDO"])
        count,backup=import_preview(self.db,rows)
        self.assertEqual(count,2); self.assertTrue(backup.exists())
        self.assertEqual(self.db.products("BAG")[0]["stock"],1000)
        self.assertEqual(self.db.products("ROLL")[0]["stock"],20)
        self.assertEqual(len(self.db.movements()),2)

    def test_existing_and_file_duplicate_sku_are_not_imported(self):
        first=preview_file(self.workbook([self.valid_bag("DUP")]),self.db); import_preview(self.db,first)
        rows=preview_file(self.workbook([self.valid_bag("DUP"),self.valid_bag("NEW"),self.valid_bag("NEW")]),self.db)
        self.assertEqual([row.status for row in rows],["DUPLICADO","VÁLIDO","DUPLICADO"])

    def test_incomplete_row_is_error(self):
        row=self.valid_bag(); row[1]=""
        preview=preview_file(self.workbook([row]),self.db)
        self.assertEqual(preview[0].status,"ERROR"); self.assertIn("nombre",preview[0].message)

    def test_non_numeric_value_is_error(self):
        row=self.valid_roll(); row[2]="treinta"
        preview=preview_file(self.workbook(roll_rows=[row]),self.db)
        self.assertEqual(preview[0].status,"ERROR"); self.assertIn("ancho_cm",preview[0].message)

    def test_csv_is_supported(self):
        path=self.root/"bags.csv"
        with open(path,"w",encoding="utf-8-sig",newline="") as handle:
            writer=csv.writer(handle); writer.writerow(BAG_COLUMNS); writer.writerow(self.valid_bag("CSV-1"))
        preview=preview_file(path,self.db)
        self.assertEqual(preview[0].status,"VÁLIDO"); self.assertEqual(preview[0].data["kind"],"BAG")

    def test_template_has_two_correct_sheets(self):
        path=create_template(self.root/"template.xlsx"); book=load_workbook(path,read_only=True)
        try:
            self.assertEqual(book.sheetnames,["Bolsas","Rollos"])
            self.assertEqual([cell.value for cell in next(book["Bolsas"].iter_rows())],BAG_COLUMNS)
            self.assertEqual([cell.value for cell in next(book["Rollos"].iter_rows())],ROLL_COLUMNS)
        finally:book.close()


if __name__=="__main__":unittest.main()
