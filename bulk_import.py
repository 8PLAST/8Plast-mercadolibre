from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


BAG_COLUMNS = ["sku_interno","nombre","ancho_cm","largo_cm","micrones","color_material",
               "tipo_categoria","pack_almacenamiento","stock_minimo","stock_objetivo","stock_inicial","activo"]
ROLL_COLUMNS = ["sku_interno","nombre","ancho_cm","micrones","metros","color_material",
                "tipo_categoria","stock_minimo","stock_objetivo","stock_inicial","activo"]


@dataclass
class PreviewRow:
    sheet: str
    row_number: int
    sku: str
    name: str
    status: str
    message: str
    data: dict | None = None


def _clean_header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def _positive_number(value, label):
    try: result=float(str(value).replace(",", "."))
    except (TypeError,ValueError): raise ValueError(f"{label}: debe ser numérico")
    if result<=0: raise ValueError(f"{label}: debe ser mayor que cero")
    return result


def _nonnegative_int(value, label):
    try:
        number=float(str(value).replace(",", "."))
        if not number.is_integer(): raise ValueError
        result=int(number)
    except (TypeError,ValueError): raise ValueError(f"{label}: debe ser un número entero")
    if result<0: raise ValueError(f"{label}: no puede ser negativo")
    return result


def _active(value):
    normalized=str(value).strip().lower()
    if normalized in ("1","si","sí","true","activo","activa","yes"): return 1
    if normalized in ("0","no","false","inactivo","inactiva"): return 0
    raise ValueError("activo: usar SI/NO, 1/0 o ACTIVO/INACTIVO")


def _validate(raw, kind):
    sku=str(raw.get("sku_interno") or "").strip(); name=str(raw.get("nombre") or "").strip()
    if not sku: raise ValueError("sku_interno: campo obligatorio")
    if not name: raise ValueError("nombre: campo obligatorio")
    minimum=_nonnegative_int(raw.get("stock_minimo"),"stock_minimo")
    target=_nonnegative_int(raw.get("stock_objetivo"),"stock_objetivo")
    initial=_nonnegative_int(raw.get("stock_inicial"),"stock_inicial")
    if target<minimum: raise ValueError("stock_objetivo: debe ser igual o mayor que stock_minimo")
    data={"kind":kind,"sku":sku,"name":name,"width_cm":_positive_number(raw.get("ancho_cm"),"ancho_cm"),
          "microns":_positive_number(raw.get("micrones"),"micrones"),
          "color_material":str(raw.get("color_material") or "").strip(),
          "category":str(raw.get("tipo_categoria") or "").strip(),"minimum_stock":minimum,
          "target_stock":target,"active":_active(raw.get("activo")),"initial_stock":initial}
    if kind=="BAG":
        data.update({"length_cm":_positive_number(raw.get("largo_cm"),"largo_cm"),"meters_per_roll":None,
                     "storage_pack":_nonnegative_int(raw.get("pack_almacenamiento"),"pack_almacenamiento")})
        if data["storage_pack"]==0: raise ValueError("pack_almacenamiento: debe ser mayor que cero")
    else:
        data.update({"length_cm":None,"meters_per_roll":_positive_number(raw.get("metros"),"metros"),"storage_pack":1})
    return data


def _rows_from_xlsx(path):
    workbook=load_workbook(path,read_only=True,data_only=True)
    try:
        for sheet_name,kind,required in (("Bolsas","BAG",BAG_COLUMNS),("Rollos","ROLL",ROLL_COLUMNS)):
            if sheet_name not in workbook.sheetnames: continue
            sheet=workbook[sheet_name]; iterator=sheet.iter_rows(values_only=True)
            headers=[_clean_header(x) for x in next(iterator,())]
            for row_number,values in enumerate(iterator,start=2):
                if not any(value not in (None,"") for value in values): continue
                yield sheet_name,row_number,kind,{headers[i]:value for i,value in enumerate(values) if i<len(headers)}
    finally: workbook.close()


def _rows_from_csv(path):
    with open(path,"r",encoding="utf-8-sig",newline="") as handle:
        sample=handle.read(4096); handle.seek(0)
        try:dialect=csv.Sniffer().sniff(sample,delimiters=",;")
        except csv.Error:dialect=csv.excel
        reader=csv.DictReader(handle,dialect=dialect); headers={_clean_header(x) for x in (reader.fieldnames or [])}
        kind="BAG" if "largo_cm" in headers else "ROLL" if "metros" in headers else ""
        for row_number,row in enumerate(reader,start=2):
            yield "CSV",row_number,kind,{_clean_header(k):v for k,v in row.items()}


def preview_file(path, database):
    path=Path(path)
    if path.suffix.lower() not in (".xlsx",".csv"): raise ValueError("Seleccioná un archivo .xlsx o .csv")
    source=_rows_from_xlsx(path) if path.suffix.lower()==".xlsx" else _rows_from_csv(path)
    existing={p["sku"].casefold() for p in database.products(active="Todos")}; seen=set(); rows=[]
    for sheet,row_number,kind,raw in source:
        sku=str(raw.get("sku_interno") or "").strip(); name=str(raw.get("nombre") or "").strip()
        if not kind:
            rows.append(PreviewRow(sheet,row_number,sku,name,"ERROR","No se pudo identificar Bolsas o Rollos")); continue
        expected=BAG_COLUMNS if kind=="BAG" else ROLL_COLUMNS
        missing_headers=[column for column in expected if column not in raw]
        if missing_headers:
            rows.append(PreviewRow(sheet,row_number,sku,name,"ERROR","Faltan columnas: "+", ".join(missing_headers))); continue
        if sku.casefold() in existing or (sku and sku.casefold() in seen):
            rows.append(PreviewRow(sheet,row_number,sku,name,"DUPLICADO","El SKU ya existe o está repetido en el archivo")); continue
        try:
            data=_validate(raw,kind); seen.add(sku.casefold())
            rows.append(PreviewRow(sheet,row_number,sku,name,"VÁLIDO","Listo para importar",data))
        except ValueError as exc: rows.append(PreviewRow(sheet,row_number,sku,name,"ERROR",str(exc)))
    if not rows: raise ValueError("El archivo no contiene filas de productos.")
    return rows


def import_preview(database, rows):
    valid=[row.data for row in rows if row.status=="VÁLIDO" and row.data]
    if not valid: raise ValueError("No hay productos válidos para importar.")
    db_path=Path(database.path); stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
    backup=db_path.with_name(f"8plast_stock_pre_carga_masiva_{stamp}.db")
    database.backup(backup)
    database.bulk_add_products(valid)
    return len(valid),backup


def create_template(path):
    workbook=Workbook(); workbook.remove(workbook.active)
    examples={"Bolsas":["BOL-5070","Bolsa 50x70 transparente",50,70,30,"Transparente","Cortada",50,500,2000,0,"SI"],
              "Rollos":["ROL-30100","Rollo bolsa tubo 30 cm",30,50,100,"Transparente","Bolsa tubo",5,20,0,"SI"]}
    for sheet_name,columns in (("Bolsas",BAG_COLUMNS),("Rollos",ROLL_COLUMNS)):
        sheet=workbook.create_sheet(sheet_name); sheet.sheet_view.showGridLines=False; sheet.freeze_panes="A2"
        sheet.append(columns); sheet.append(examples[sheet_name]); sheet.auto_filter.ref=f"A1:{sheet.cell(1,len(columns)).coordinate}"
        for cell in sheet[1]:
            cell.font=Font(name="Segoe UI",bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="2673DD"); cell.alignment=Alignment(horizontal="center")
        for index,column in enumerate(columns,start=1):
            width=max(12,min(28,len(column)+3)); sheet.column_dimensions[sheet.cell(1,index).column_letter].width=width
        active_col=columns.index("activo")+1; validation=DataValidation(type="list",formula1='"SI,NO"',allow_blank=False)
        sheet.add_data_validation(validation); validation.add(f"{sheet.cell(2,active_col).coordinate}:{sheet.cell(500,active_col).coordinate}")
    workbook.save(path)
    return Path(path)
